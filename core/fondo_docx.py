"""Schema di regolamento e di bando per il fondo comunale energia (Word)."""
import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _e(x):
    return (f"{x:,.0f} €".replace(",", ".")) if x else "[●] €"


def documento(comune: str, modello: str, opz: dict, misure: list[dict]) -> bytes:
    """modello: 'perduto' o 'rotativo'. misure: [{'misura', 'unita', 'massimale'}]."""
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    comune = comune or "[●]"
    rot = modello == "rotativo"
    nome_fondo = opz.get("nome") or "Fondo comunale energia per i cittadini"

    h = d.add_heading(f"COMUNE DI {comune.upper()}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(f"SCHEMA DI REGOLAMENTO DEL «{nome_fondo.upper()}»").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(f"Bozza di lavoro generata il {date.today().strftime('%d/%m/%Y')} con lo strumento «Compensazioni FER · FVG». "
                        "Da verificare con il Segretario comunale e il Responsabile del servizio finanziario.")
    p.runs[0].italic = True

    art = [
        ("Oggetto e finalità",
         f"Il presente regolamento disciplina il {nome_fondo} (il «Fondo»), destinato a sostenere interventi di efficienza "
         "energetica e di diffusione delle fonti rinnovabili nelle abitazioni dei residenti, con priorità per le famiglie in "
         "condizione di povertà energetica. Il Fondo attua il programma di misure di compensazione ambientale e territoriale "
         "previsto dalla convenzione stipulata con [●] in data [●], ai sensi del DM 10 settembre 2010, Allegato 2, "
         "del D.Lgs. 190/2024 e della L.R. FVG 2/2025."),
        ("Risorse",
         "Il Fondo è alimentato dalle somme versate a titolo di compensazione, iscritte in bilancio come entrata a destinazione "
         "vincolata, e da eventuali ulteriori risorse comunali, regionali o di terzi. Le somme non utilizzate nell'anno restano "
         "vincolate alle finalità del Fondo."
         + (f" Dotazione iniziale prevista: {_e(opz.get('capitale_iniziale'))}." if opz.get("capitale_iniziale") else "")
         + (f" Quota annua prevista: {_e(opz.get('quota_annua'))}." if opz.get("quota_annua") else "")),
        ("Beneficiari",
         "Possono accedere le persone fisiche residenti nel Comune, per interventi sull'abitazione di residenza situata nel "
         "territorio comunale."
         + (f" È data priorità ai nuclei con ISEE non superiore a {_e(opz.get('isee_priorita'))}." if opz.get("isee_priorita") else "")),
        ("Interventi ammessi",
         "Sono ammessi esclusivamente gli interventi elencati nell'Allegato A, con i relativi massimali. La Giunta può aggiornare "
         "l'elenco ogni anno, nel rispetto delle finalità di cui all'art. 1."),
    ]
    if rot:
        art.append(("Forma dell'agevolazione",
                    "Il Fondo opera come fondo rotativo: concede, tramite il soggetto gestore convenzionato (istituto bancario o "
                    f"confidi), finanziamenti a tasso zero di importo non superiore a {_e(opz.get('prestito_max'))}, rimborsabili in "
                    f"{opz.get('durata_rimborso') or '[●]'} anni con rate [mensili/annuali]. I rientri alimentano nuovamente il Fondo. "
                    "Il soggetto gestore istruisce le domande, eroga i finanziamenti, cura la riscossione e rendiconta al Comune; "
                    "i rapporti sono regolati da apposita convenzione."))
    else:
        art.append(("Forma dell'agevolazione",
                    f"Il Fondo concede contributi a fondo perduto pari al {opz.get('perc_contributo') or '[●]'}% della spesa ammissibile, "
                    f"entro i massimali dell'Allegato A e comunque non oltre {_e(opz.get('tetto_famiglia'))} per nucleo familiare per anno."))
    art += [
        ("Cumulabilità",
         "Il contributo [è / non è] cumulabile con altre agevolazioni pubbliche per il medesimo intervento (detrazioni fiscali, "
         "Conto Termico, bonus nazionali), nei limiti della spesa effettivamente sostenuta. Il richiedente dichiara le altre "
         "agevolazioni richieste o ottenute."),
        ("Bando annuale",
         "Ogni anno il Comune approva un bando che indica risorse disponibili, termini, modalità di presentazione e criteri di "
         "graduatoria, nel rispetto del presente regolamento (art. 12 della L. 241/1990)."),
        ("Requisiti tecnici",
         "Gli apparecchi devono rispettare le classi energetiche minime indicate nell'Allegato A. "
         + ("È obbligatoria la rottamazione dell'apparecchio sostituito, documentata dal rivenditore. " if opz.get("rottamazione") else "")
         + "Gli interventi edilizi rispettano le norme vigenti."),
        ("Controlli e revoca",
         "Il Comune effettua controlli a campione su almeno il [10]% delle pratiche. In caso di dichiarazioni non veritiere o di "
         "interventi non conformi, l'agevolazione è revocata e le somme sono recuperate."),
        ("Rendicontazione e trasparenza",
         "Ogni anno il Comune pubblica sul proprio sito un rendiconto con risorse entrate, agevolazioni concesse, numero di nuclei "
         "raggiunti e risparmi energetici stimati, e lo trasmette al comitato di monitoraggio della convenzione, se costituito."),
        ("Trattamento dei dati", "I dati sono trattati ai sensi del Regolamento UE 2016/679."),
        ("Entrata in vigore", "Il regolamento entra in vigore dalla data di esecutività della deliberazione di approvazione."),
    ]
    for i, (t, x) in enumerate(art, 1):
        d.add_heading(f"Art. {i} – {t}", level=2)
        d.add_paragraph(x)

    d.add_heading("Allegato A – Interventi ammessi e massimali", level=2)
    tb = d.add_table(rows=1, cols=3)
    tb.style = "Light Grid Accent 1"
    tb.rows[0].cells[0].text, tb.rows[0].cells[1].text, tb.rows[0].cells[2].text = "Intervento", "Unità", "Massimale"
    for m in misure or [{"misura": "[●]", "unita": "", "massimale": 0}]:
        c = tb.add_row().cells
        c[0].text, c[1].text, c[2].text = m["misura"], m.get("unita", ""), _e(m.get("massimale"))

    d.add_page_break()
    h = d.add_heading(f"SCHEMA DI BANDO {date.today().year + 1} – {nome_fondo.upper()}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t, x in [
        ("Risorse disponibili", f"{_e(opz.get('quota_annua'))} per l'anno [●]."),
        ("Chi può partecipare", "Persone fisiche residenti nel Comune, per l'abitazione di residenza."),
        ("Interventi finanziabili", "Quelli dell'Allegato A del regolamento, con i massimali indicati."),
        ("Agevolazione", ("Finanziamento a tasso zero tramite il soggetto gestore [●]." if rot else
                          f"Contributo pari al {opz.get('perc_contributo') or '[●]'}% della spesa, fino a {_e(opz.get('tetto_famiglia'))} per nucleo.")),
        ("Graduatoria", "Punteggi: ISEE [●] punti; presenza di persone fragili [●]; risparmio energetico atteso [●]; ordine di arrivo a parità."),
        ("Domanda", "Dal [●] al [●], tramite [sportello / PEC / piattaforma], con preventivo o fattura, dichiarazione ISEE e documentazione tecnica."),
        ("Erogazione", "Entro [●] giorni dalla presentazione della fattura quietanzata e della documentazione di rottamazione, se richiesta."),
        ("Informazioni", "Sportello energia del Comune, [orari, contatti]."),
    ]:
        d.add_heading(t, level=2)
        d.add_paragraph(x)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
