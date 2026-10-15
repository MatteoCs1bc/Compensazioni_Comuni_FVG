"""Modelli per l'impresa: proposta di compensazione, sollecito, atto unilaterale d'obbligo."""
import io
from datetime import date

from docx import Document
from docx.shared import Pt


def _e(x):
    return (f"{x:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")) if x else "[●] €"


def _doc():
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    return d


def _intestazione(d, p, oggetto):
    d.add_paragraph(f"{p.get('proponente') or '[Società]'}\n[indirizzo, PEC]")
    d.add_paragraph(f"Spett.le Comune di {p.get('comune') or '[●]'}\nAlla c.a. del Sindaco e del Responsabile del procedimento\n"
                    "PEC: [●]\n\ne p.c. Regione Autonoma Friuli Venezia Giulia – Servizio transizione energetica\nPEC: [●]")
    d.add_paragraph(f"[Luogo], {date.today().strftime('%d/%m/%Y')}")
    par = d.add_paragraph()
    par.add_run(f"Oggetto: {oggetto}").bold = True


def _riepilogo(p):
    s = (f"impianto «{p.get('nome') or '[●]'}», potenza {p.get('potenza_mw') or '[●]'} MW"
         + (f" con accumulo da {p['bess_mw']:g} MW" if p.get("bess_mw") else "")
         + f", nel Comune di {p.get('comune') or '[●]'}")
    return s


def _tabella_interventi(d, interventi):
    if not interventi:
        d.add_paragraph("[Elenco degli interventi proposti, da concordare con il Comune]")
        return
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Intervento", "Quantità", "Importo stimato"
    for x in interventi:
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = x["nome"], x["quantita"], _e(x["costo"])


def proposta(p: dict, procedura: str, interventi: list[dict]) -> bytes:
    d = _doc()
    _intestazione(d, p, f"Proposta di programma di compensazioni ambientali e territoriali – {_riepilogo(p)}")
    d.add_paragraph(
        f"Con riferimento al progetto in oggetto, per il quale è in corso / sarà avviata la procedura di {procedura}, la scrivente "
        "Società trasmette la propria proposta di programma di compensazioni ambientali e territoriali a favore del Comune, ai "
        "sensi del D.Lgs. 190/2024 (artt. 8 e 9), della L.R. FVG 2/2025 (art. 5, c. 1, lett. i) e del DM 10 settembre 2010, Allegato 2.")
    d.add_heading("Valore e metodo di calcolo", level=2)
    d.add_paragraph(f"Valore complessivo proposto: {_e(p.get('compensazione_eur'))}"
                    + (f", pari a {_e(p.get('compensazione_annua_eur'))} all'anno" if p.get("compensazione_annua_eur") else "")
                    + ". " + (p.get("calcolo_testo") or "[Metodo di calcolo]"))
    d.add_heading("Interventi proposti", level=2)
    _tabella_interventi(d, interventi)
    d.add_heading("Richiesta", level=2)
    d.add_paragraph(
        "Si chiede al Comune di comunicare entro 30 giorni dal ricevimento della presente le proprie osservazioni e la disponibilità "
        "a definire la convenzione, di cui si allega uno schema. La Società è disponibile a un incontro e a una presentazione "
        "pubblica del progetto ai cittadini (L.R. FVG 2/2025, art. 5, c. 1, lett. j).")
    d.add_paragraph("Allegati: schema di convenzione; calcolo delle compensazioni; elaborati di progetto.")
    d.add_paragraph("\nDistinti saluti\n[Firma del legale rappresentante]")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def sollecito(p: dict, data_prima: str) -> bytes:
    d = _doc()
    _intestazione(d, p, f"Sollecito – proposta di compensazioni del {data_prima or '[●]'} – {_riepilogo(p)}")
    d.add_paragraph(
        f"Con PEC del {data_prima or '[●]'} la scrivente Società ha trasmesso al Comune la proposta di programma di compensazioni "
        "per il progetto in oggetto, chiedendo un riscontro entro 30 giorni. A oggi non risulta pervenuta alcuna risposta.")
    d.add_paragraph(
        "Si rinnova la richiesta di un riscontro entro ulteriori 15 giorni. In mancanza, la Società provvederà a formalizzare i "
        "propri impegni con atto unilaterale d'obbligo, che sarà depositato presso il Comune e trasmesso all'amministrazione "
        "procedente, ferma restando la piena disponibilità a definire una convenzione condivisa.")
    d.add_paragraph("La presente è trasmessa per conoscenza alla Regione, amministrazione procedente.")
    d.add_paragraph("\nDistinti saluti\n[Firma del legale rappresentante]")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def atto_unilaterale(p: dict, procedura: str, interventi: list[dict], garanzia: bool = True) -> bytes:
    d = _doc()
    d.add_heading("ATTO UNILATERALE D'OBBLIGO", level=1)
    d.add_paragraph("per la realizzazione delle misure di compensazione ambientale e territoriale")
    par = d.add_paragraph("Bozza di lavoro generata con lo strumento «Compensazioni FER · FVG»: da far verificare a un legale. "
                          "Se va autenticato da un notaio o registrato dipende dall'uso che se ne fa.")
    par.runs[0].italic = True
    d.add_paragraph(f"La società {p.get('proponente') or '[●]'}, con sede in [●], C.F./P.IVA [●], in persona del legale "
                    "rappresentante [●] (la «Società»),")
    d.add_heading("premesso che", level=2)
    for x in [
        f"la Società ha presentato / è titolare del titolo di {procedura} per l'{_riepilogo(p)} (l'«Impianto»);",
        "la normativa (D.Lgs. 190/2024, artt. 8 e 9; L.R. FVG 2/2025, art. 5, c. 1, lett. i; DM 10/9/2010, All. 2) prevede "
        "misure di compensazione a favore dei Comuni interessati;",
        "con PEC del [●] e del [●] la Società ha proposto al Comune un programma di compensazioni e uno schema di convenzione, "
        "senza ottenere riscontro;",
        "la Società intende comunque rendere certi e verificabili i propri impegni;",
    ]:
        d.add_paragraph(x, style="List Bullet")
    d.add_heading("si obbliga", level=2)
    obblighi = [
        f"a destinare a misure di compensazione a favore del Comune di {p.get('comune') or '[●]'} un importo complessivo di "
        f"{_e(p.get('compensazione_eur'))}" + (f", pari a {_e(p.get('compensazione_annua_eur'))} all'anno per la durata di esercizio "
                                               "dell'Impianto" if p.get("compensazione_annua_eur") else "") + ";",
        "a realizzare o finanziare gli interventi dell'allegato, o altri interventi delle medesime categorie indicati dal Comune "
        "entro [●] giorni dalla richiesta;",
        "a rivalutare annualmente gli importi secondo l'indice ISTAT FOI;",
        "a trasmettere al Comune entro il 31 marzo di ogni anno il rendiconto delle somme destinate e degli interventi realizzati;",
        "a trasferire gli obblighi del presente atto a chi dovesse subentrare nella titolarità dell'Impianto o del titolo;",
        "a sottoscrivere in qualsiasi momento una convenzione con il Comune che recepisca, anche migliorandoli, i presenti impegni.",
    ]
    if garanzia:
        obblighi.insert(3, "a prestare, prima dell'inizio dei lavori, fideiussione bancaria o assicurativa a prima richiesta di "
                           "importo pari a [●] € a garanzia dei presenti obblighi;")
    for i, x in enumerate(obblighi, 1):
        d.add_paragraph(f"{i}. {x}")
    d.add_paragraph("Il presente atto è depositato presso il Comune e trasmesso all'amministrazione procedente.")
    d.add_heading("Allegato – Interventi", level=2)
    _tabella_interventi(d, interventi)
    d.add_paragraph("\n[Luogo e data]\n[Firma del legale rappresentante]")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()
