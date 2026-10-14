"""Estrae dai documenti di progetto (PDF, DOCX, TXT) i dati utili al calcolo.

Non è infallibile: propone valori che l'utente conferma o corregge.
"""
from __future__ import annotations

import io
import re

PAT_POT = re.compile(r"(\d{1,3}(?:[.\s]\d{3})*(?:,\d+)?|\d+(?:[.,]\d+)?)\s*(MWp|MWe|MW|kWp|kWe|kW)\b", re.I)
CONTESTO_POT = re.compile(r"potenza|nominale|complessiv|picco|immissione|installat", re.I)


def testo_da_file(nome: str, dati: bytes) -> str:
    n = nome.lower()
    if n.endswith(".pdf"):
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(dati))
        return "\n".join((p.extract_text() or "") for p in r.pages[:80])
    if n.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(dati))
        parti = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parti.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parti)
    return dati.decode("utf-8", errors="ignore")


def _num(s: str) -> float:
    s = s.replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") >= 1 and len(s.split(".")[-1]) == 3:
        s = s.replace(".", "")
    return float(s)


def potenze(testo: str) -> list[tuple[float, str]]:
    """Potenze in MW trovate vicino a parole come 'potenza', con il contesto."""
    out = []
    for m in PAT_POT.finditer(testo):
        ctx = testo[max(0, m.start() - 120): m.end() + 20]
        if not CONTESTO_POT.search(ctx):
            continue
        try:
            v = _num(m.group(1))
        except ValueError:
            continue
        unita = m.group(2).lower()
        mw = v / 1000 if unita.startswith("k") else v
        if 0.005 <= mw <= 2000:
            out.append((round(mw, 4), " ".join(ctx.split())))
    return out


def tipo(testo: str) -> str | None:
    t = testo.lower()
    conti = {
        "agrivoltaico": len(re.findall(r"agrivoltaic|agro-?voltaic|agrofotovoltaic", t)),
        "eolico": len(re.findall(r"\beolic|aerogenerator|turbin[ae] eolic", t)),
        "bess": len(re.findall(r"\bbess\b|accumulo elettrochimico|sistema di accumulo|storage", t)),
        "fv_flottante": len(re.findall(r"flottant", t)),
        "fv_terra": len(re.findall(r"fotovoltaic", t)),
    }
    if conti["fv_flottante"] > 2:
        return "fv_flottante"
    if conti["agrivoltaico"] > 2:
        return "agrivoltaico"
    if conti["eolico"] > 3:
        return "eolico"
    if conti["bess"] > conti["fv_terra"] and conti["bess"] > 2:
        return "bess"
    if conti["fv_terra"]:
        return "fv_terra"
    return None


def comuni_citati(testo: str, elenco: list[str]) -> list[str]:
    t = testo.lower()
    trovati = []
    for c in elenco:
        k = c.lower()
        n = len(re.findall(r"\b" + re.escape(k) + r"\b", t))
        if n:
            trovati.append((n, c))
    return [c for _, c in sorted(trovati, reverse=True)]


def proponente(testo: str) -> str | None:
    m = re.search(r"([A-Z][A-Za-z0-9&'.\- ]{2,60}?\s(?:S\.?\s?r\.?\s?l\.?|S\.?\s?p\.?\s?A\.?|SRL|SPA))\b", testo)
    return " ".join(m.group(1).split()) if m else None


def procedura(testo: str) -> str | None:
    t = testo.lower()
    if "autorizzazione unica" in t:
        return "au"
    if "procedura abilitativa semplificata" in t or re.search(r"\bpas\b", t):
        return "pas"
    return None


def analizza(testo: str, elenco_comuni: list[str]) -> dict:
    pot = potenze(testo)
    scelta = None
    if pot:
        def punteggio(item):
            v, ctx = item
            c = ctx.lower()
            s = 0
            s += 3 if "complessiv" in c else 0
            s += 2 if "nominale" in c else 0
            s += 1 if ("picco" in c or "installat" in c) else 0
            s -= 4 if re.search(r"modul|inverter|cabina|trasformator|gruppo di conversione|ciascun", c) else 0
            return (s, v)
        scelta = max(pot, key=punteggio)[0]
    return {
        "potenza_mw": scelta,
        "potenze_trovate": pot[:12],
        "tipo": tipo(testo),
        "comuni": comuni_citati(testo, elenco_comuni)[:5],
        "proponente": proponente(testo),
        "procedura": procedura(testo),
        "via": bool(re.search(r"valutazione di impatto ambientale|assoggettabilit", testo, re.I)),
        "caratteri": len(testo),
    }
