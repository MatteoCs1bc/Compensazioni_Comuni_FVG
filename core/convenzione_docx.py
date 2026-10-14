"""Genera una bozza di convenzione in formato Word, precompilata con i dati del progetto."""
import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _euro(x):
    return (f"{x:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")) if x else "[●] €"


def bozza(p: dict, opzioni: dict, interventi: list[dict]) -> bytes:
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    comune = p.get("comune") or "[●]"
    prop = p.get("proponente") or "[●]"
    nome = p.get("nome") or "[●]"

    t = d.add_heading("CONVENZIONE PER LA REALIZZAZIONE DELLE MISURE DI COMPENSAZIONE AMBIENTALE E TERRITORIALE", level=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("BOZZA DI LAVORO generata con lo strumento «Compensazioni FER · FVG» il "
                    f"{date.today().strftime('%d/%m/%Y')}. Da verificare con il Segretario comunale e gli uffici competenti.").italic = True

    d.add_paragraph("TRA")
    d.add_paragraph(f"il COMUNE DI {comune.upper()}, rappresentato da [●], in qualità di [●], autorizzato con deliberazione "
                    "della Giunta comunale n. [●] del [●] (il «Comune»)")
    d.add_paragraph(f"E\n{prop}, con sede in [●], C.F./P.IVA [●], rappresentata da [●] (il «Proponente»)")

    d.add_heading("Premesse", level=2)
    prem = [
        f"il Proponente intende realizzare l'impianto «{nome}» di potenza pari a {p.get('potenza_mw') or '[●]'} MW"
        + (f", con sistema di accumulo da {p['bess_mw']:g} MW" if p.get("bess_mw") else "")
        + f", nel territorio del Comune di {comune} (l'«Impianto»);",
        f"per l'Impianto è in corso / è stato rilasciato il titolo [{opzioni.get('titolo', 'autorizzazione unica / PAS')}] n. [●] del [●];",
        "l'art. 8, c. 4, lett. m), n. 2 e l'art. 9, c. 10, lett. d) del D.Lgs. 190/2024 prevedono un programma di compensazioni "
        "ambientali e territoriali a favore dei Comuni interessati;",
        "l'art. 5, c. 1, lett. i) della L.R. FVG 2/2025 prevede, per gli impianti soggetti ad autorizzazione unica sopra 1 MW, "
        "compensazioni non inferiori al 3 per cento dei proventi;",
        "il DM 10 settembre 2010, Allegato 2, stabilisce che le misure compensative non hanno carattere meramente patrimoniale e "
        "sono orientate a interventi di miglioramento ambientale, efficienza energetica, diffusione delle fonti rinnovabili e "
        "sensibilizzazione della cittadinanza;",
    ]
    for x in prem:
        d.add_paragraph(x, style="List Bullet")

    d.add_heading("Art. 1 – Oggetto", level=2)
    d.add_paragraph("La presente convenzione disciplina le misure di compensazione a favore del Comune, i tempi, le modalità di "
                    "realizzazione e le garanzie.")

    d.add_heading("Art. 2 – Valore delle compensazioni", level=2)
    perc = p.get("percentuale")
    d.add_paragraph(
        f"Il valore complessivo delle compensazioni è pari a {_euro(p.get('compensazione_eur'))}"
        + (f", corrispondente al {perc:.1f}% dei proventi stimati dell'Impianto".replace(".", ",") if perc else "")
        + (f", di cui {_euro(p.get('compensazione_annua_eur'))} all'anno" if p.get("compensazione_annua_eur") else "")
        + ". Il calcolo è riportato nell'Allegato 2."
    )
    if opzioni.get("annuale"):
        d.add_paragraph("Le compensazioni sono corrisposte con cadenza annuale per tutta la durata di esercizio dell'Impianto, "
                        "a partire dalla data di entrata in esercizio.")
    if opzioni.get("rivalutazione"):
        d.add_paragraph("Gli importi sono rivalutati annualmente secondo l'indice ISTAT FOI.")

    d.add_heading("Art. 3 – Interventi", level=2)
    d.add_paragraph("Le compensazioni sono destinate agli interventi elencati nell'Allegato 1, che le Parti possono aggiornare di "
                    "comune accordo. " + ("Gli interventi sono realizzati dal Comune; il Proponente versa le somme a fronte della "
                    "rendicontazione." if opzioni.get("realizza") == "Comune" else
                    "Gli interventi sono realizzati dal Proponente, previa approvazione dei progetti da parte del Comune."))

    if opzioni.get("garanzia"):
        d.add_heading("Art. 4 – Garanzie", level=2)
        d.add_paragraph("A garanzia degli obblighi della presente convenzione il Proponente presta, prima dell'inizio dei lavori, "
                        "fideiussione bancaria o assicurativa a prima richiesta di importo pari a [●] €, distinta dalle garanzie "
                        "per la dismissione dell'Impianto.")
    d.add_heading("Art. 5 – Durata ed efficacia", level=2)
    d.add_paragraph("La convenzione è sottoscritta prima dell'inizio dei lavori e resta efficace per tutta la vita dell'Impianto, "
                    "fino al completo ripristino dei luoghi.")
    if opzioni.get("subentro"):
        d.add_heading("Art. 6 – Subentro", level=2)
        d.add_paragraph("In caso di cessione dell'Impianto, del titolo autorizzativo o di quote della società, gli obblighi della "
                        "presente convenzione si trasferiscono al subentrante, che li assume espressamente.")
    if opzioni.get("rendiconto"):
        d.add_heading("Art. 7 – Monitoraggio e rendiconto", level=2)
        d.add_paragraph("Entro il 31 marzo di ogni anno il Proponente trasmette al Comune un rendiconto delle somme versate e degli "
                        "interventi realizzati, che il Comune pubblica sul proprio sito.")
    if opzioni.get("penali"):
        d.add_heading("Art. 8 – Penali", level=2)
        d.add_paragraph("Per ogni mese di ritardo nei versamenti o nella realizzazione degli interventi si applica una penale pari "
                        "a [●] €, salvo il maggior danno.")
    d.add_heading("Art. 9 – Controversie", level=2)
    d.add_paragraph("Le Parti tentano una composizione bonaria; in mancanza è competente il Foro di [●].")

    d.add_heading("Allegato 1 – Interventi", level=2)
    if interventi:
        tb = d.add_table(rows=1, cols=3)
        tb.style = "Light Grid Accent 1"
        h = tb.rows[0].cells
        h[0].text, h[1].text, h[2].text = "Intervento", "Quantità", "Importo stimato"
        for x in interventi:
            c = tb.add_row().cells
            c[0].text, c[1].text, c[2].text = x["nome"], x["quantita"], _euro(x["costo"])
    else:
        d.add_paragraph("[Elenco degli interventi]")
    d.add_heading("Allegato 2 – Calcolo delle compensazioni", level=2)
    d.add_paragraph(opzioni.get("calcolo") or "[Inserire il calcolo]")
    d.add_paragraph("\nPer il Comune ______________________          Per il Proponente ______________________")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
